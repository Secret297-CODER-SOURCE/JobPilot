"""Извлечение текста из резюме разных форматов (pdf/docx/txt)."""


def extract_text(file_path, fmt):
    fmt = fmt.lower()
    if fmt == "pdf":
        return _extract_pdf(file_path)
    if fmt == "docx":
        return _extract_docx(file_path)
    if fmt == "txt":
        return _extract_txt(file_path)
    raise ValueError(f"Неподдерживаемый формат: {fmt}")


def _extract_pdf(file_path):
    from pypdf import PdfReader
    reader = PdfReader(str(file_path))
    parts = []
    for page in reader.pages:
        text = page.extract_text(extraction_mode="layout") or page.extract_text() or ""
        if text:
            parts.append(text)
    return "\n".join(parts).strip()


def _extract_docx(file_path):
    import docx
    doc = docx.Document(str(file_path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts).strip()


def _extract_txt(file_path):
    with open(file_path, "rb") as f:
        raw = f.read()
    for encoding in ("utf-8", "cp1251", "latin-1"):
        try:
            return raw.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace").strip()
